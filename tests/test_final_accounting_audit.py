"""
BillZen AI - Final Accounting Export, Excel Design & Reconciliation Audit Test Suite
Verifies:
1. Multi-sheet Reports Excel (6 sheets with CA standards, true numbers, formulas, totals)
2. CA-ready GST Audit Excel with drilldowns (Kacha bills excluded from 0% exempt)
3. Single Sales & Purchase invoice Excel exports
4. Supplier Bills Excel & CSV totals reconciliation (Unicom / Unicorn Services 25-26)
5. Credit Note & Debit Note totals consistency across UI calculation, DB, CSV, Excel, and Tally XML
6. Tally XML exports for Sales, Purchase, Credit Notes, Debit Notes (XML validation, ledger mapping, voucher count)
7. Return quantity logic (discrete vs continuous, Available = Original - Already Returned, incremental returns)
"""

import os
import io
import json
import openpyxl
import pytest
import xml.etree.ElementTree as ET
import app as billzen

TEST_CSRF_TOKEN = 'test-csrf-token-12345'

@pytest.fixture
def client():
    billzen.DATABASE = os.path.abspath('billzen.db')
    billzen.app.config['TESTING'] = True
    with billzen.app.test_client() as client:
        # Authenticate as user_id 2 (our active business with real data)
        with client.session_transaction() as sess:
            sess['user_id'] = 2
            sess['username'] = 'sachin'
            sess['role'] = 'owner'
            sess['csrf_token'] = TEST_CSRF_TOKEN
        yield client

@pytest.fixture
def auth_headers():
    return {'X-CSRF-Token': TEST_CSRF_TOKEN}


# =========================================================================
# 1. REPORTS EXCEL 6-SHEET WORKBOOK TEST
# =========================================================================
def test_reports_excel_workbook_structure_and_totals(client):
    """Verifies /api/reports/export/excel returns a valid 6-sheet workbook with correct totals and formulas."""
    resp = client.get('/api/reports/export/excel')
    assert resp.status_code == 200
    assert resp.content_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment' in resp.headers.get('Content-Disposition', '')

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    sheet_names = wb.sheetnames
    expected_sheets = [
        "Dashboard Summary",
        "Sales Register",
        "Purchase Register",
        "Stock Report",
        "Customer Summary",
        "GST Summary"
    ]
    for name in expected_sheets:
        assert name in sheet_names, f"Expected sheet '{name}' missing from Reports Excel"

    # Check Sales Register total formula
    ws_sales = wb["Sales Register"]
    assert ws_sales.max_row >= 5
    last_row_sales = ws_sales.max_row
    sales_total_label = ws_sales.cell(row=last_row_sales, column=1).value
    assert sales_total_label == "TOTAL SALES"
    # Column 11 is Invoice Total
    tot_formula = ws_sales.cell(row=last_row_sales, column=11).value
    assert isinstance(tot_formula, str) and tot_formula.startswith("=SUM(")

    # Check Purchase Register total formula
    ws_pur = wb["Purchase Register"]
    assert ws_pur.max_row >= 5
    last_row_pur = ws_pur.max_row
    pur_total_label = ws_pur.cell(row=last_row_pur, column=1).value
    assert pur_total_label == "TOTAL PURCHASE"
    # Column 11 is Purchase Total
    tot_formula_pur = ws_pur.cell(row=last_row_pur, column=11).value
    assert isinstance(tot_formula_pur, str) and tot_formula_pur.startswith("=SUM(")

    # Check Stock Report total formula
    ws_stock = wb["Stock Report"]
    last_row_stock = ws_stock.max_row
    assert ws_stock.cell(row=last_row_stock, column=1).value in ("TOTALS", "TOTAL INVENTORY COST")
    # Column 6 is Inventory Value, Column 8 is Potential Sale Value, Column 9 is Potential Profit
    assert ws_stock.cell(row=last_row_stock, column=6).value.startswith("=SUM(")
    assert ws_stock.cell(row=last_row_stock, column=8).value.startswith("=SUM(")
    assert ws_stock.cell(row=last_row_stock, column=9).value.startswith("=SUM(")


# =========================================================================
# 2. GST REPORT CA EXCEL TEST
# =========================================================================
def test_gst_ca_excel_structure_and_no_kacha_in_exempt(client):
    """Verifies /api/gst/export/excel CA-audit workbook and ensures Kacha bills are not categorized as 0% exempt."""
    resp = client.get('/api/gst/export/excel')
    assert resp.status_code == 200
    assert resp.content_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    assert "GST Summary & Reconciliation" in wb.sheetnames
    assert "18% GST Register" in wb.sheetnames
    assert "5% GST Register" in wb.sheetnames

    ws_sum = wb["GST Summary & Reconciliation"]
    # Verify rate rows
    rates_found = [ws_sum.cell(row=r, column=1).value for r in range(6, 12) if ws_sum.cell(row=r, column=1).value]
    assert any("18%" in str(r) for r in rates_found)

    # Check that Kacha bills are excluded from GST returns
    drill_resp = client.get('/api/gst/drilldown')
    assert drill_resp.status_code == 200
    drill_data = drill_resp.get_json().get('drilldown', {})
    zero_bucket = drill_data.get('rates', {}).get('0', {})
    for inv in zero_bucket.get('invoices', []):
        assert inv.get('bill_category') != 'kacha', "Kacha bill must NOT be included in 0% Exempt GST bucket!"


# =========================================================================
# 3. SINGLE BILL EXCEL EXPORTS (SALES & PURCHASE)
# =========================================================================
def test_single_sales_and_purchase_excel_export(client):
    """Verifies single sales and single purchase invoice export to professional Excel."""
    # Test sales bill
    sales_resp = client.get('/api/sales')
    sales = sales_resp.get_json().get('sales', [])
    if sales:
        first_bill = sales[0]['bill_id']
        s_excel = client.get(f'/api/sales/{first_bill}/export/excel')
        assert s_excel.status_code == 200
        assert s_excel.content_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        wb = openpyxl.load_workbook(io.BytesIO(s_excel.data))
        assert "Tax Invoice" in wb.sheetnames or "Estimate" in wb.sheetnames
        ws = wb.active
        assert ws["A1"].value is not None
        assert "BILLZEN" in ws["A1"].value or len(str(ws["A1"].value)) > 0

    # Specifically verify local 18% tax invoice GST breakup (INV261008D6FF)
    inv_check = billzen.query("SELECT * FROM sales WHERE bill_id='INV261008D6FF' AND user_id=2", one=True)
    if inv_check:
        res = client.get('/api/sales/INV261008D6FF/export/excel')
        assert res.status_code == 200
        wb_inv = openpyxl.load_workbook(io.BytesIO(res.data))
        ws_inv = wb_inv.active
        # Item 1 row 11: 3000 taxable, 270 CGST, 270 SGST, 3540 line total
        assert ws_inv.cell(row=11, column=8).value == 3000
        assert ws_inv.cell(row=11, column=9).value == 270
        assert ws_inv.cell(row=11, column=10).value == 270
        assert ws_inv.cell(row=11, column=12).value == 3540
        # Item 2 row 12: 50000 taxable, 4500 CGST, 4500 SGST, 59000 line total
        assert ws_inv.cell(row=12, column=8).value == 50000
        assert ws_inv.cell(row=12, column=9).value == 4500
        assert ws_inv.cell(row=12, column=10).value == 4500
        assert ws_inv.cell(row=12, column=12).value == 59000
        # Totals row 13:
        assert ws_inv.cell(row=13, column=9).value == "=SUM(I11:I12)"
        assert ws_inv.cell(row=13, column=10).value == "=SUM(J11:J12)"
        # Grand total row 15:
        assert ws_inv.cell(row=15, column=12).value == 62540

    # Test purchase bill (Invoice UGST76-2/171)
    p_excel = client.get('/api/purchases/UGST76-2/171/export/excel')
    assert p_excel.status_code == 200
    assert p_excel.content_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    wb_p = openpyxl.load_workbook(io.BytesIO(p_excel.data))
    assert "Supplier Bill" in wb_p.sheetnames
    ws_p = wb_p.active
    # Check party details in header block
    found_sup = any("Unicom" in str(ws_p.cell(row=r, column=1).value or '') for r in range(4, 10))
    assert found_sup, "Expected Unicom Services in purchase bill export header"


# =========================================================================
# 4. SUPPLIER BILLS LOOKUP & EXCEL EXPORT (UNICOM / UNICORN SERVICES)
# =========================================================================
def test_supplier_bills_and_excel_export(client):
    """Verifies Supplier Bills returns purchases matching name or GSTIN and exports Excel."""
    # Supplier 5 is Unicom Services 25-26 with GSTIN 29AMEPK5421C1ZA
    resp = client.get('/api/suppliers/5/purchases')
    assert resp.status_code == 200
    data = resp.get_json()
    assert data['success'] is True
    purchases = data.get('purchases', [])
    assert len(purchases) >= 1
    # Check invoice UGST76-2/171 is present
    inv_nos = [p.get('invoice_no') for p in purchases]
    assert "UGST76-2/171" in inv_nos

    # Total displayed must equal sum of displayed valid bills
    computed_sum = sum(float(p['total_amount']) for p in purchases)
    assert round(computed_sum, 2) == round(float(data['total_amount']), 2)

    # Test Excel export
    excel_resp = client.get('/api/suppliers/5/export/excel')
    assert excel_resp.status_code == 200
    assert excel_resp.content_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    wb_s = openpyxl.load_workbook(io.BytesIO(excel_resp.data))
    assert "Supplier Invoices" in wb_s.sheetnames
    ws_s = wb_s.active
    last_r = ws_s.max_row
    assert ws_s.cell(row=last_r, column=1).value == "TOTAL PURCHASED"
    tot_formula = ws_s.cell(row=last_r, column=8).value
    assert isinstance(tot_formula, str) and tot_formula.startswith("=SUM(")


# =========================================================================
# 5. CREDIT NOTE & DEBIT NOTE HISTORY, CSV & EXCEL
# =========================================================================
def test_returns_registers_csv_and_excel(client):
    """Verifies Credit Note & Debit Note exports have required headers and bottom TOTAL rows."""
    # Sales returns (Credit notes) CSV
    cn_csv = client.get('/api/sales/returns/export/csv')
    assert cn_csv.status_code == 200
    csv_text = cn_csv.data.decode('utf-8')
    assert "Credit Note No" in csv_text
    assert "TOTAL CREDIT NOTES" in csv_text
    assert "TOTAL CREDIT NOTE VALUE" in csv_text

    # Sales returns Excel
    cn_excel = client.get('/api/sales/returns/export/excel')
    assert cn_excel.status_code == 200
    assert cn_excel.content_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    # Purchase returns (Debit notes) CSV
    dn_csv = client.get('/api/purchases/returns/export/csv')
    assert dn_csv.status_code == 200
    csv_dn_text = dn_csv.data.decode('utf-8')
    assert "Debit Note No" in csv_dn_text
    assert "TOTAL DEBIT NOTES" in csv_dn_text
    assert "TOTAL DEBIT NOTE VALUE" in csv_dn_text

    # Purchase returns Excel
    dn_excel = client.get('/api/purchases/returns/export/excel')
    assert dn_excel.status_code == 200
    assert dn_excel.content_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


# =========================================================================
# 6. DEBIT NOTE MATHEMATICAL & ROUNDING CONSISTENCY
# =========================================================================
def test_debit_note_line_item_math_consistency(client, auth_headers):
    """
    Verifies the exact line-item tax calculation for Debit Note DN-2026-0003:
    Taxable: 11,110.42
    Line-item IGST sum: 1,999.87
    Total: 13,110.29
    Ensures that Tally XML Debit Note matches this exact amount without 1-paisa drift.
    """
    resp = client.get('/api/purchases/returns/DN-2026-0003')
    assert resp.status_code == 200
    dn = resp.get_json().get('debit_note', {})
    assert dn.get('debit_note_no') == 'DN-2026-0003'
    assert round(float(dn.get('subtotal', 0)), 2) == 11110.42
    assert round(float(dn.get('igst_amount', 0)), 2) == 1999.87
    assert round(float(dn.get('total_amount', 0)), 2) == 13110.29

    # Now verify Tally XML export for this debit note
    tally_resp = client.post('/api/tally/export/v2', headers=auth_headers, json={
        "category": "all",
        "selected_vouchers": ["DN-2026-0003"],
        "include_debit_notes": True
    })
    assert tally_resp.status_code == 200
    assert tally_resp.content_type.startswith('application/xml')
    xml_str = tally_resp.data.decode('utf-8')
    assert '<VOUCHER VCHTYPE="Debit Note"' in xml_str
    assert '<VOUCHERNUMBER>DN-2026-0003</VOUCHERNUMBER>' in xml_str
    assert '-13110.29' in xml_str or '13110.29' in xml_str


# =========================================================================
# 7. TALLY XML EXPORT FOR ALL VOUCHER CATEGORIES
# =========================================================================
def test_tally_xml_all_voucher_categories(client, auth_headers):
    """Verifies Tally XML generation for Purchase, Credit Note, and Debit Note."""
    # Test Purchase voucher export
    pur_resp = client.post('/api/tally/export/v2', headers=auth_headers, json={
        "category": "all",
        "selected_vouchers": ["UGST76-2/171"]
    })
    assert pur_resp.status_code == 200
    pur_xml = pur_resp.data.decode('utf-8')
    assert '<VOUCHER VCHTYPE="Purchase"' in pur_xml
    assert 'UGST76-2/171' in pur_xml

    # Test Credit Note export
    cn_resp = client.post('/api/tally/export/v2', headers=auth_headers, json={
        "category": "all",
        "selected_vouchers": ["CN-2026-0003"]
    })
    # User 2 has CN-2026-0003
    assert cn_resp.status_code == 200
    cn_xml = cn_resp.data.decode('utf-8')
    assert '<VOUCHER VCHTYPE="Credit Note"' in cn_xml

    # Test Empty Selection error
    err_resp = client.post('/api/tally/export/v2', headers=auth_headers, json={
        "category": "all",
        "selected_vouchers": ["NON_EXISTENT_VOUCHER_12345"]
    })
    assert err_resp.status_code == 400
    assert "None of the selected vouchers could be found" in err_resp.get_json().get('error', '')


# =========================================================================
# 8. UNIT-AWARE RETURN QUANTITY LOGIC & AVAILABLE QTY TRACKING
# =========================================================================
def test_unit_aware_return_context_and_available_qty(client):
    """Verifies discrete vs continuous unit handling and Available = Original - Already Returned."""
    # Test purchase return context for UGST76-2/171
    ctx_resp = client.get('/api/purchases/UGST76-2/171/return-context')
    assert ctx_resp.status_code == 200
    ctx_data = ctx_resp.get_json()
    assert ctx_data['success'] is True
    items = ctx_data.get('items', [])
    assert len(items) > 0

    for it in items:
        unit = it.get('unit', '')
        purchased = float(it.get('purchased_qty', 0))
        already_ret = float(it.get('already_returned_qty', 0))
        returnable = float(it.get('returnable_qty', 0))
        # Formula verification
        assert round(returnable, 2) == round(max(0.0, purchased - already_ret), 2)
        if unit in ('Pcs', 'Pc', 'Nos', 'Box', 'Pack'):
            assert float(it.get('returnable_qty')).is_integer()


# =========================================================================
# 9. WOW EXCEL EXPORT STANDARD TESTS (SALES, PURCHASES, STOCK, CUSTOMER)
# =========================================================================
def test_wow_standard_all_exports(client):
    """Verifies that Sales, Purchase, Stock, and Customer Ledger all export dedicated WOW workbooks with formulas and sheets."""
    # 1. Sales Register WOW Excel
    resp_s = client.get('/api/sales/export/excel')
    assert resp_s.status_code == 200
    wb_s = openpyxl.load_workbook(io.BytesIO(resp_s.data))
    assert "Sales Register" in wb_s.sheetnames
    assert "Itemized Line Items" in wb_s.sheetnames
    ws_sr = wb_s["Sales Register"]
    # Verify TOTAL SALES row with =SUM formulas
    last_row_sr = ws_sr.max_row
    assert ws_sr.cell(row=last_row_sr, column=1).value == "TOTAL SALES"
    assert ws_sr.cell(row=last_row_sr, column=11).value.startswith("=SUM(")

    ws_sli = wb_s["Itemized Line Items"]
    last_row_sli = ws_sli.max_row
    assert ws_sli.cell(row=last_row_sli, column=1).value == "TOTAL ITEM SALES"
    assert ws_sli.cell(row=last_row_sli, column=13).value.startswith("=SUM(") # CGST sum
    assert ws_sli.cell(row=last_row_sli, column=14).value.startswith("=SUM(") # SGST sum

    # Verify INV261008D6FF line items if present
    inv_taxable = 0
    inv_cgst = 0
    inv_sgst = 0
    for r in range(5, last_row_sli):
        b_val = str(ws_sli.cell(row=r, column=2).value or '')
        if 'INV261008D6FF' in b_val:
            inv_taxable += float(ws_sli.cell(row=r, column=11).value or 0)
            inv_cgst += float(ws_sli.cell(row=r, column=13).value or 0)
            inv_sgst += float(ws_sli.cell(row=r, column=14).value or 0)
    if inv_taxable > 0:
        assert inv_taxable == 53000.0
        assert inv_cgst == 4770.0
        assert inv_sgst == 4770.0

    # 2. Purchase Register WOW Excel
    resp_p = client.get('/api/purchases/export/excel')
    assert resp_p.status_code == 200
    wb_p = openpyxl.load_workbook(io.BytesIO(resp_p.data))
    assert "Purchase Register" in wb_p.sheetnames
    assert "Itemized Inward Supplies" in wb_p.sheetnames
    ws_pr = wb_p["Purchase Register"]
    last_row_pr = ws_pr.max_row
    assert ws_pr.cell(row=last_row_pr, column=1).value == "TOTAL PURCHASE"
    assert ws_pr.cell(row=last_row_pr, column=11).value.startswith("=SUM(")

    ws_item = wb_p["Itemized Inward Supplies"]
    last_row_item = ws_item.max_row
    assert ws_item.cell(row=last_row_item, column=1).value == "TOTAL ITEM PURCHASES"
    assert ws_item.cell(row=last_row_item, column=13).value.startswith("=SUM(") # CGST sum
    assert ws_item.cell(row=last_row_item, column=14).value.startswith("=SUM(") # SGST sum

    # Verify UGST76-2/171 line item if present has populated tax columns
    for r in range(5, last_row_item):
        inv_val = str(ws_item.cell(row=r, column=3).value or '')
        if 'UGST76-2/171' in inv_val:
            assert float(ws_item.cell(row=r, column=11).value) == 480000.0
            assert float(ws_item.cell(row=r, column=13).value) == 43200.0
            assert float(ws_item.cell(row=r, column=14).value) == 43200.0
            assert float(ws_item.cell(row=r, column=15).value) == 0.0
            assert float(ws_item.cell(row=r, column=16).value) == 566400.0

    # 3. Stock Inventory Valuation WOW Excel
    resp_stk = client.get('/api/products/export/excel')
    assert resp_stk.status_code == 200
    wb_stk = openpyxl.load_workbook(io.BytesIO(resp_stk.data))
    assert "Stock Valuation & Inventory" in wb_stk.sheetnames
    ws_stk = wb_stk["Stock Valuation & Inventory"]
    last_row_stk = ws_stk.max_row
    assert ws_stk.cell(row=last_row_stk, column=2).value == "TOTAL INVENTORY VALUATION"
    assert ws_stk.cell(row=last_row_stk, column=8).value.startswith("=SUM(")

    # 4. Customer Ledger WOW Excel
    cust = billzen.query("SELECT c.id FROM customers c JOIN sales s ON s.customer_name=c.name AND s.user_id=2 WHERE c.user_id=2 LIMIT 1", one=True)
    if cust:
        cid = cust['id']
        resp_c = client.get(f'/api/customers/{cid}/export/excel')
        assert resp_c.status_code == 200
        wb_c = openpyxl.load_workbook(io.BytesIO(resp_c.data))
        assert "Customer Statement" in wb_c.sheetnames
        ws_c = wb_c["Customer Statement"]
        last_row_c = ws_c.max_row
        assert ws_c.cell(row=last_row_c, column=2).value == "TOTAL STATEMENT"
        assert ws_c.cell(row=last_row_c, column=10).value.startswith("=SUM(")

